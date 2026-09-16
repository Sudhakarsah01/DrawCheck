import streamlit as st
import os, re, base64
import pandas as pd
from datetime import datetime
from PIL import Image
from io import BytesIO
from collections import Counter
import pymupdf as fitz
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="DrawCheck", page_icon="🏗️", layout="wide")

# ── Brand ─────────────────────────────────────────────────
PRIMARY="#3D4BA0"; PRIMARY_D="#2C3680"; PRIMARY_L="#E8EAF6"
ACCENT="#5C6BC0"; WHITE="#FFFFFF"; GRAY_L="#F0F2F8"; GRAY_M="#E0E4F0"
TEXT_D="#1A1A2E"; TEXT_M="#4A4A6A"
RED="#D32F2F"; AMBER="#F57C00"; GREEN="#388E3C"

st.markdown(f"""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&family=Poppins:wght@600;700&display=swap');
html,body,[class*="css"]{{font-family:'Inter',sans-serif;color:{TEXT_D};}}
#MainMenu{{visibility:hidden;}}header{{visibility:hidden;}}footer{{visibility:hidden;}}
.stApp{{background:{WHITE};}}
.hbar{{background:linear-gradient(135deg,{PRIMARY_D} 0%,{PRIMARY} 60%,{ACCENT} 100%);
       margin:-1rem -1rem 2rem -1rem;border-radius:0 0 16px 16px;
       box-shadow:0 4px 20px rgba(61,75,160,.25);}}
.hinner{{display:flex;align-items:center;justify-content:space-between;padding:18px 36px;}}
.htitle{{text-align:center;flex:1;padding:0 24px;}}
.htitle h1{{font-family:'Poppins',sans-serif;font-size:26px;font-weight:700;color:{WHITE};margin:0;}}
.htitle p{{font-size:13px;color:rgba(255,255,255,.80);margin:3px 0 0;font-weight:300;}}
.hbadge{{background:rgba(255,255,255,.15);border:1px solid rgba(255,255,255,.30);
         border-radius:20px;padding:6px 16px;color:{WHITE};font-size:12px;font-weight:500;}}
.mrow{{display:flex;gap:14px;margin:16px 0;flex-wrap:wrap;}}
.mcard{{flex:1;min-width:120px;background:{WHITE};border:1px solid {GRAY_M};
        border-radius:14px;padding:18px 16px;text-align:center;
        box-shadow:0 2px 8px rgba(61,75,160,.08);}}
.mcard.blue{{border-top:4px solid {PRIMARY};}}
.mcard.red{{border-top:4px solid {RED};}}
.mcard.amber{{border-top:4px solid {AMBER};}}
.mcard.green{{border-top:4px solid {GREEN};}}
.mcard.gray{{border-top:4px solid #9E9E9E;}}
.mval{{font-family:'Poppins',sans-serif;font-size:32px;font-weight:700;color:{TEXT_D};line-height:1;margin-bottom:6px;}}
.mcard.red .mval{{color:{RED};}}.mcard.amber .mval{{color:{AMBER};}}.mcard.green .mval{{color:{GREEN};}}
.mlbl{{font-size:12px;color:{TEXT_M};font-weight:500;text-transform:uppercase;letter-spacing:.05em;}}
.sh{{font-family:'Poppins',sans-serif;font-size:16px;font-weight:600;color:{PRIMARY};
     border-left:4px solid {PRIMARY};padding-left:12px;margin:24px 0 12px;}}
.spanel{{background:{GRAY_L};border-radius:14px;padding:18px;border:1px solid {GRAY_M};}}
.spanel h4{{font-family:'Poppins',sans-serif;color:{PRIMARY};font-size:14px;font-weight:600;margin:0 0 12px;}}
.fcard{{background:{GRAY_L};border-radius:14px;padding:20px;text-align:center;border:1px solid {GRAY_M};}}
.ficon{{font-size:32px;margin-bottom:10px;}}
.ftitle{{font-family:'Poppins',sans-serif;font-size:14px;font-weight:600;color:{PRIMARY};margin-bottom:6px;}}
.fdesc{{font-size:12px;color:{TEXT_M};line-height:1.5;}}
.stButton>button{{background:linear-gradient(135deg,{PRIMARY} 0%,{ACCENT} 100%);
  color:{WHITE};border:none;border-radius:10px;font-weight:600;font-size:14px;
  padding:10px 20px;transition:all .2s;box-shadow:0 3px 10px rgba(61,75,160,.25);}}
.stDownloadButton>button{{background:linear-gradient(135deg,{GREEN} 0%,#43A047 100%) !important;
  color:{WHITE} !important;border:none !important;border-radius:10px !important;
  font-weight:600 !important;font-size:15px !important;padding:12px 24px !important;
  box-shadow:0 4px 14px rgba(56,142,60,.30) !important;}}
hr{{border:none;border-top:1px solid {GRAY_M};margin:20px 0;}}
</style>""", unsafe_allow_html=True)

def img_b64(path):
    try:
        with open(path,"rb") as f: return base64.b64encode(f.read()).decode()
    except: return ""

BASE = os.path.dirname(os.path.abspath(__file__))
logo = img_b64(os.path.join(BASE,"logo_transparent.png"))
logo_html = f'<img src="data:image/png;base64,{logo}" style="height:50px"/>' if logo else \
            '<span style="color:white;font-weight:700;font-size:18px">SG Design Nepal</span>'







st.markdown(f"""
<div class="hbar"><div class="hinner">
  <div>{logo_html}</div>
  <div class="htitle"><h1>DrawCheck</h1>
    <p>HVAC Drawing QA Annotation Extractor — SG Design Nepal</p></div>
  <div class="hbadge">v1.0 Professional</div>
</div></div>""", unsafe_allow_html=True)



# ════════════════════════════════════════════════════════════
# EXTRACTION ENGINE
# ════════════════════════════════════════════════════════════
HVAC_PREFIXES=[
    "FCU-","AHU-","FD-","VCD-","MVCD-","EAF-","OAF-","AP-","BMS-",
    "HUM-","SP-","MSD-","CHWP-","HWP-","BT-","ERAC-","MSSB-","IFD-",
    "DAF-","EAD-","SAD-","NRD-","ORC-","TRC-","ERC-","CU-","CDU-",
    "RCU-","BS-","SS-","RL-","SL-","TL-","EL-","DPM-","SFD-","EFD-",
    "ACU-","PAU-","MAU-","ERV-","HRV-","RTU-","SAF-","RAF-","BC-",
]
NOISE_CONTENT={"f/a","t/a","f/b","s/a","r/a","e/a","o/a","c/w","n/a","tbc","tbd","nts"}
NOISE_PAT=[
    r'^\d+[\d,]*\.?\d*\s*(mm|m|sq\.?m|l/s|ls|kg|kpa|pa|%)?$',
    r'^\d+[xX]\d+([xX]\d+)?$',
    r'^[øØoO]?\d+[øØ]?\s*[xX]?\s*\d*$',
    r'^([øØ]?\d+[øØ]?\s*){2,}$',
    r'^\(?\d+\.?\d*\s*[lL]/[sS]\)?$',
    r'^\d+\s*[lL]/[sS]\s+[A-Z]{2,5}-',
    r'^[A-Z]{2,6}-[A-Z0-9]+-[A-Z0-9]+-[A-Z0-9]+-\d+$',
    r'^([A-Z]{2,5}-[A-Z0-9]+-[A-Z0-9]+-\d+\s*){2,}$',
    r'^S\d+-[A-Z]-\d+(\s+\d+)?$',
    r'^\d+\s*-\s*[A-Z]$',
    r'^\d+[xX]\d+\s+\d+$',
    r'^[Cc]/[Ww]\s+\d+[xX]\d+\s+\d+[xX]\d+',
    r'^[A-Z]{1,2}\d*$',r'^P\d{2,3}$',
    r'^\d{1,2}/\d{2}/\d{2,4}$',r'^[A-Z]{3}\s*[\']\d{2}$',
    r'^\d{1,2}:\d{2}(:\d{2})?\s*[AP]M$',
    r'^[A-Z0-9]{5,}-[A-Z0-9]{3,}-[A-Z0-9]{3,}-',
    r'^\d{1,2}\s*/\s*\d{1,2}$',r'^@\d+%',
    r'^[A-Z]{1,3}\d+-[A-Z]{1,2}-\d+',r'^\d+\.\s+ALL\s+',
    r'^(Level\s*)?\d+$',r'^(GF|B\d+|L\d+|RF|UG|P\d)$',
    r'^[\d\s\.\,\/\(\)mmMM]+$',
]
LEGEND_NOISE=[
    "uninsulated metal ductwork","internally insulated ductwork",
    "externally insulated ductwork","insulated flexible ductwork",
    "uninsulated flexible ductwork","fire rated ductwork",
    "supply air ductwork","exhaust air ductwork","return air ductwork",
    "outside air ductwork","transfer air ductwork","toilet exhaust",
    "kitchen exhaust ductwork","ceiling access panel","duct access panel",
    "temperature sensor","fire damper","volume control damper",
    "acoustic flexible ductwork","natural ventilation",
    "digital touchscreen panel","wall mounted thermostat",
    "room controller","co2 sensor","all duct sizes are metal",
    "all ducts are flange joint","all louvres sizes are neck",
    "security clips to be documented","all openings contain ember",
    "all return air grilles to be complete","model used",
    "checked by auditor","for construction","preliminary issue",
    "cc3 issue","for information","project no","project number",
    "drawing no","drawing number","drawing title","revision",
    "description","issued by","client","contractor","builder",
    "key plan","links","services",
]
QA_KW=[
    "missing","wrong","check","confirm","add","remove","write","change",
    "fix","provide","show","move","replace","correct","update","ensure",
    "verify","blanked","spelling","incorrect","error","not required",
    "instead","coordinate","clash","overlap","not able","not sufficient",
    "non compliant","fail","incomplete","hyphen not","not matching",
    "confirm if","make sure","double check","louvre required",
    "bulkhead required","penetration required","use","place","put",
    "refer","install","route","relocate","resize","flip","align",
    "avoid","section","view","3d","please","why","what","where","how",
    "maybe","might","could","would","want","don't","doesn't","shouldn't",
    "can't","cannot","is not","are not","is this","isn't","how will",
    "can we","should this","which one","are these","is there",
    "tag missing","tag not","tag clash","grille tag","duct tag",
    "ap missing","access panel","duct ap","ceiling ap","internal lining",
    "lining missing","full height","pipe connection","spigot",
    "open ended","close ended","return air path","filter",
    "insulation","sensor missing","nrd direction","bou","bot","cl ",
    "setdown","set down","return air","supply air","exhaust air",
    "outside air","fire rated","fire rating","smoke","compartment",
    "clash with","pass through","minimum","maximum","required",
    "small letter","capital letter","spelling mistake","spelling error",
    "arrow missing","symbol not correct","overlap","overlapping",
    "too close","too sharp","smooth radius","shift","rotate","hide",
    "layout","draw flex","draw duct","draw pipe","properly",
    "make up","doorway","open doorway","makeup via","typical comment",
    "please confirm","please check","please show","please provide",
    "please add","please move","please use","please remove",
    "confirm this","confirm if","double confirm","section required",
    "3d view","better to provide","better to show","change to",
    "louvre required","bulkhead required","penetration required",
    "minimum gap","distance between","min.","max.","off ",
    "fall towards","toward","interval","radius","reducer",
    "sharp","extend","lower","raise","increase","reduce",
    "incomplete pipe","spigot access","add ap","add 1 ap",
    "ap for this","spigot & ap","ap & tag","tag & arrow",
    "rotate cu","shift cu","shift bulkhead","shift it","move it",
    "hide layout","text overlap","text overlapping",
]

HIGH_KW=["missing","wrong","incorrect","error","must","critical","remove",
         "clash","clashing","not required","full height","cannot","not able",
         "non compliant","fail","incomplete","hyphen not","not matching",
         "louvre required","bulkhead required","penetration required",
         "symbol not correct","spelling mistake","spelling error",
         "too sharp","too close","incomplete pipe","no ","never","urgent"]
MED_KW=["check","confirm","provide","show","move","update","change","add",
        "ensure","verify","review","consider","please","should","recommend",
        "suggest","clarify","query","avoid","need to","why","better to",
        "confirm if","make sure","double check","rotate","shift","reduce",
        "increase","hide","layout","is there","can we","how will",
        "minimum","maximum","which one","isn't","isn't this"]
CAT_RULES=[
    (["tag missing","label missing","no tag","missing tag","grille tag",
      "duct tag","louvre tag","hyphen not","tag clash","small letter",
      "capital letter","symbol not correct","spelling mistake",
      "wrong nrd","nrd direction","arrow missing","tag & arrow"],"Tag Issue"),
    (["ap missing","access panel","ap clashing","duct ap","ap not",
      "ceiling ap","duct mounted ap","flip ap","600x600 ap",
      "not able to access","add ap","add 1 ap","ap for this",
      "spigot & ap","ap & tag"],"Access Panel Issue"),
    (["insulation missing","internal lining","lining missing",
      "not fire rated","insulation lines","dashed","insulated",
      "fire rated","fire rating"],"Insulation Issue"),
    (["clash","clashing","conflict","interference","full height",
      "opening to pass","pass through","run through","don't allow",
      "clash with slab"],"Clash Issue"),
    (["duct","flex","spigot","fitting","plenum","transition","pipe",
      "oval","setdown","set down","bou","bot","top of","offset",
      "route","reroute","reducer","radius","fall towards","smooth radius",
      "draw flex","draw duct","sharp","too sharp","200 dia"],"Duct Issue"),
    (["fcu","ahu","fan coil","equipment","coil","grille","louvre",
      "nrd","sensor","co sensor","co2","mvcd","ifd","door grille",
      "make up","doorway","rotate cu","shift cu"],"Equipment Issue"),
    (["dimension","level","height","distance","cl ","centre line",
      "minimum","maximum","min.","max.","gap required","interval",
      "not full height"],"Dimension Issue"),
    (["fire","fd","fire damper","compartment","smoke","fire rated",
      "not fire rated","fire egress","egress"],"Fire Issue"),
    (["section","view","3d","detail","legend","note","label",
      "ga plan","room tag","drawing","hide layout","text overlap",
      "overlapping","spelling"],"Drawing Issue"),
    (["riser","vertical","shaft","upper levels","can't these risers"],"Riser Issue"),
    (["schedule","model","type","product","not matching",
      "open ended plenum","close ended plenum","plenum box",
      "confirm which","which one is"],"Schedule Issue"),
    (["return air","r/a path","return air path"],"Return Air Issue"),
    (["louvre required","bulkhead required","penetration required",
      "make up","open doorway","makeup via","natural ventilation",
      "filter missing","filter in o/a"],"Ventilation Issue"),
    (["incomplete pipe","pipe connection","spigot access"],"Pipe Issue"),
]

def _classify(text):
    if not text: return "General Issue","Low"
    t=text.lower()
    sev="High" if any(k in t for k in HIGH_KW) else \
        "Medium" if any(k in t for k in MED_KW) else "Low"
    cat="General Issue"
    for kws,category in CAT_RULES:
        if any(k in t for k in kws): cat=category; break
    return cat,sev



def _is_noise(text):
    t=text.strip(); tl=t.lower().strip()
    if len(t)<3: return True
    if tl in NOISE_CONTENT: return True
    if any(leg in tl for leg in LEGEND_NOISE): return True
    for pat in NOISE_PAT:
        if re.match(pat,t,re.IGNORECASE): return True
    return False

def _is_qa(text):
    tl=text.lower()
    return any(kw in tl for kw in QA_KW)



def _nearest_tag(cx,cy,words_pos,radius=600):
    best_tag,best_dist="",radius
    for wx,wy,word in words_pos:
        wup=word.upper().rstrip(".,;)([]")
        if any(wup.startswith(p.upper()) for p in HVAC_PREFIXES):
            dist=((wx-cx)**2+(wy-cy)**2)**0.5
            if dist<best_dist: best_dist=dist; best_tag=word.rstrip(".,;)([]")
    return best_tag



def _loc(cx,cy,pw,ph):
    nx,ny=cx/pw,cy/ph
    v="Top" if ny<0.20 else "Upper" if ny<0.40 else "Middle" if ny<0.60 else "Lower" if ny<0.80 else "Bottom"
    h="Left" if nx<0.20 else "Left-Ctr" if nx<0.40 else "Centre" if nx<0.60 else "Right-Ctr" if nx<0.80 else "Right"
    return f"{v}-{h}"

def _norm(text):
    return re.sub(r'\s+',' ',text.lower().strip())[:70]

def _dwg_info(page):
    ph=page.rect.height
    info={"drawing_number":"","revision":""}
    full=""
    for block in page.get_text("blocks"):
        if block[1]>ph*0.80: full+=" "+block[4].strip()
    m=re.search(r'[A-Z]{2,6}-[A-Z0-9]+-[A-Z0-9]+-[A-Z]+-[A-Z0-9]+',full)
    if m: info["drawing_number"]=m.group(0)
    m=re.search(r'\bRev\s*([A-Z0-9]{1,3})\b',full,re.I)
    if m: info["revision"]=m.group(1)
    return info

def _tb_y(page):
    ph=page.rect.height; pw=page.rect.width; tb_y=ph
    for block in page.get_text("blocks"):
        btext=block[4].lower().strip(); by0=block[1]; bx0=block[0]
        if any(m in btext for m in ["drawing no","project no","checked by","issued by","for construction"]):
            if by0>ph*0.75 and bx0>pw*0.50:
                if by0<tb_y: tb_y=by0
    return max(tb_y,ph*0.80)




def extract_all_markups(pdf_bytes):
    doc=fitz.open(stream=pdf_bytes,filetype="pdf")
    results={}
    for page_num in range(len(doc)):
        page=doc[page_num]
        pw,ph=page.rect.width,page.rect.height
        tb_y=_tb_y(page)
        words_pos=[]
        for w in page.get_text("words"):
            words_pos.append(((w[0]+w[2])/2,(w[1]+w[3])/2,w[4]))
        dwg_info=_dwg_info(page)
        markups=[]; seen=set(); idx=1
        for annot in page.annots():
            atype=annot.type[1]
            if atype not in ("FreeText","Text"): continue
            info=annot.info
            content=" ".join(info.get("content","").split()).strip()
            subject=" ".join(info.get("subject","").split()).strip()
            title=info.get("title","").strip()
            text=content or subject
            if not text or len(text)<3: continue
            rect=annot.rect
            cx,cy=(rect.x0+rect.x1)/2,(rect.y0+rect.y1)/2
            in_tb=(cy>tb_y) and (cx>pw*0.55)
            if in_tb: continue
            if _is_noise(text): continue
            if not _is_qa(text): continue
            norm=_norm(text)
            if norm in seen: continue
            if any(norm[:50] in e for e in seen): continue
            seen.add(norm)
            cat,sev=_classify(text)
            tag=_nearest_tag(cx,cy,words_pos)
            loc=_loc(cx,cy,pw,ph)
            markups.append({"id":idx,"page":page_num+1,
                "drawing_number":dwg_info["drawing_number"],
                "location":loc,"comment":text,"affected_tag":tag,
                "category":cat,"severity":sev,"reviewer":title})
            idx+=1
        results[page_num+1]={"markups":markups,"drawing_info":dwg_info,"total":len(markups)}
    doc.close()
    return results

# ════════════════════════════════════════════════════════════
# EXCEL BUILDER
# ════════════════════════════════════════════════════════════
XN,XB,XL="1F3864","3D4BA0","E8EAF6"
XW="FFFFFF"; XRB,XRF="FFEBEE","C62828"; XYB,XYF="FFF3E0","E65100"
XGB,XGF="E8F5E9","2E7D32"; XA="F5F6FC"



def xf(h): return PatternFill("solid",fgColor=h)
def xft(bold=False,color="000000",size=10):
    return Font(bold=bold,color=color,size=size,name="Calibri")
def xal(h="left",wrap=False):
    return Alignment(horizontal=h,vertical="center",wrap_text=wrap)
def xbd():
    s=Side(style="thin",color="C5CAE9")
    return Border(left=s,right=s,top=s,bottom=s)

def xc(ws,r,c,v,bold=False,fg="000000",bg=None,align="left",
       wrap=False,size=10,mt=None):
    if mt: ws.merge_cells(f"{get_column_letter(c)}{r}:{get_column_letter(mt)}{r}")
    x=ws.cell(row=r,column=c,value=v)
    x.font=xft(bold,fg,size); x.alignment=xal(align,wrap); x.border=xbd()
    if bg: x.fill=xf(bg)
    return x

def xs(ws,r,c,sev):
    s={"High":(XRB,XRF),"Medium":(XYB,XYF),"Low":(XGB,XGF)}
    bg,fg=s.get(sev,("FFFFFF","000000"))
    x=ws.cell(row=r,column=c,value=sev)
    x.font=xft(bold=True,color=fg); x.fill=xf(bg)
    x.alignment=xal("center"); x.border=xbd()

def build_summary(wb,all_results,filename):
    ws=wb.active; ws.title="SUMMARY"
    for col,w in zip("ABCDEFGH",[6,28,44,10,10,10,14,12]):
        ws.column_dimensions[col].width=w
    ws.merge_cells("A1:H1")
    t=ws["A1"]; t.value="DRAWCHECK AI — COMPLETE QA ANNOTATION REPORT"
    t.font=Font(bold=True,color=XW,size=15,name="Calibri")
    t.fill=xf(XN); t.alignment=xal("center"); ws.row_dimensions[1].height=32
    ws.merge_cells("A2:H2")
    d=ws["A2"]
    d.value=f"File: {filename}   |   Generated: {datetime.now().strftime('%d/%m/%Y %H:%M')}   |   SG Design Nepal"
    d.font=Font(color="3D4BA0",size=10,name="Calibri",bold=True)
    d.fill=xf(XL); d.alignment=xal("center"); ws.row_dimensions[2].height=18
    all_m=[m for r in all_results.values() for m in r["markups"]]
    total=len(all_m)
    high=sum(1 for m in all_m if m["severity"]=="High")
    med=sum(1 for m in all_m if m["severity"]=="Medium")
    lo=sum(1 for m in all_m if m["severity"]=="Low")
    sheets=[p for p,r in all_results.items() if r["total"]>0]
    xc(ws,3,1,f"TOTAL: {total}",bold=True,fg=XW,bg=XB,size=13,align="center",mt=2)
    xc(ws,3,3,f"HIGH: {high}",bold=True,fg=XRF,bg=XRB,size=12,align="center",mt=4)
    xc(ws,3,5,f"MEDIUM: {med}",bold=True,fg=XYF,bg=XYB,size=12,align="center",mt=6)
    xc(ws,3,7,f"LOW: {lo}",bold=True,fg=XGF,bg=XGB,size=12,align="center")
    xc(ws,3,8,f"SHEETS: {len(sheets)}",bold=True,fg=XW,bg=XB,align="center",size=12)
    ws.row_dimensions[3].height=26
    xc(ws,5,1,"PER SHEET BREAKDOWN",bold=True,fg=XW,bg=XN,align="center",size=11,mt=8)
    ws.row_dimensions[5].height=20
    for ci,h in enumerate(["Sheet","Drawing No","Total","High","Medium","Low","Reviewer","Tab"],1):
        xc(ws,6,ci,h,bold=True,fg="3D4BA0",bg=XL,align="center")
    ws.row_dimensions[6].height=18
    ri=7
    for pg in sorted(sheets):
        r=all_results[pg]; mkps=r["markups"]
        dwg=r["drawing_info"].get("drawing_number","") or f"Page {pg}"
        h_=sum(1 for m in mkps if m["severity"]=="High")
        me_=sum(1 for m in mkps if m["severity"]=="Medium")
        lo_=sum(1 for m in mkps if m["severity"]=="Low")
        rvr=mkps[0]["reviewer"] if mkps else ""
        alt=XA if ri%2==0 else XW
        xc(ws,ri,1,pg,align="center",bg=alt,bold=True)
        xc(ws,ri,2,dwg[:38],bg=alt,size=9)
        xc(ws,ri,3,len(mkps),align="center",bg=alt,bold=True)
        for ci_,val,rbg,rfg in [(4,h_,XRB,XRF),(5,me_,XYB,XYF),(6,lo_,XGB,XGF)]:
            x=ws.cell(row=ri,column=ci_,value=val)
            x.font=xft(bold=True,color=rfg); x.fill=xf(rbg if val>0 else alt)
            x.alignment=xal("center"); x.border=xbd()
        xc(ws,ri,7,rvr,bg=alt,size=9)
        xc(ws,ri,8,f"Sheet {pg}",bg=alt,size=9,align="center")
        ws.row_dimensions[ri].height=16; ri+=1
    ri+=1
    xc(ws,ri,1,"CATEGORY BREAKDOWN",bold=True,fg=XW,bg=XN,align="center",size=11,mt=8)
    ws.row_dimensions[ri].height=20; ri+=1
    xc(ws,ri,1,"#",bold=True,fg="3D4BA0",bg=XL,align="center")
    xc(ws,ri,2,"Category",bold=True,fg="3D4BA0",bg=XL,mt=6)
    xc(ws,ri,7,"Count",bold=True,fg="3D4BA0",bg=XL,align="center",mt=8)
    ws.row_dimensions[ri].height=18; ri+=1
    for i,(cat,count) in enumerate(Counter(m["category"] for m in all_m).most_common(),1):
        alt=XA if i%2==0 else XW
        xc(ws,ri,1,i,align="center",bg=alt)
        xc(ws,ri,2,cat,bg=alt,mt=6)
        xc(ws,ri,7,count,align="center",bg=alt,bold=True,mt=8)
        ws.row_dimensions[ri].height=16; ri+=1
    ws.freeze_panes="A4"



def build_master_tab(wb,all_results):
    ws=wb.create_sheet("ALL MARKUPS")
    for col,w in zip("ABCDEFGH",[5,7,13,48,22,22,10,14]):
        ws.column_dimensions[col].width=w
    ws.merge_cells("A1:H1")
    t=ws["A1"]; t.value="ALL QA MARKUPS — COMPLETE DRAWING SET"
    t.font=Font(bold=True,color=XW,size=14,name="Calibri")
    t.fill=xf(XN); t.alignment=xal("center"); ws.row_dimensions[1].height=28
    for ci,h in enumerate(["No.","Sheet","Location","QA Markup Comment",
                            "Affected Tag","Category","Sev.","Reviewer"],1):
        xc(ws,2,ci,h,bold=True,fg="3D4BA0",bg=XL,align="center")
    ws.row_dimensions[2].height=18; ws.freeze_panes="A3"
    gidx=1
    for pg in sorted(all_results.keys()):
        for m in all_results[pg]["markups"]:
            alt=XA if gidx%2==0 else XW
            xc(ws,gidx+2,1,gidx,align="center",bg=alt)
            xc(ws,gidx+2,2,m["page"],align="center",bg=alt,bold=True)
            xc(ws,gidx+2,3,m["location"],bg=alt,size=9)
            xc(ws,gidx+2,4,m["comment"],bg=alt,wrap=True,size=9)
            xc(ws,gidx+2,5,m["affected_tag"],align="center",bg=alt,size=9)
            xc(ws,gidx+2,6,m["category"],bg=alt,size=9)
            xs(ws,gidx+2,7,m["severity"])
            xc(ws,gidx+2,8,m["reviewer"],bg=alt,size=9)
            ws.row_dimensions[gidx+2].height=40 if len(m["comment"])>80 else 20
            gidx+=1

def build_page_tab(wb,sheet_name,markups,drawing_info):
    ws=wb.create_sheet(sheet_name)
    for col,w in zip("ABCDEFGH",[5,13,48,22,22,10,14,10]):
        ws.column_dimensions[col].width=w
    ws.merge_cells("A1:H1")
    t=ws["A1"]; t.value=f"QA MARKUP ANNOTATIONS — {sheet_name}"
    t.font=Font(bold=True,color=XW,size=13,name="Calibri")
    t.fill=xf(XN); t.alignment=xal("center"); ws.row_dimensions[1].height=26
    dwg=drawing_info.get("drawing_number","") or sheet_name
    rev=drawing_info.get("revision","") or "—"
    h_=sum(1 for m in markups if m["severity"]=="High")
    me_=sum(1 for m in markups if m["severity"]=="Medium")
    lo_=sum(1 for m in markups if m["severity"]=="Low")
    ws.merge_cells("A2:C2")
    xc(ws,2,1,f"Drawing: {dwg}",bold=True,fg=XW,bg=XB)
    xc(ws,2,4,f"Rev: {rev}",bold=True,fg=XW,bg=XB,align="center")
    xc(ws,2,5,f"H:{h_}  M:{me_}  L:{lo_}",bold=True,fg=XW,bg=XB,align="center")
    ws.merge_cells("F2:G2")
    xc(ws,2,6,f"Total: {len(markups)}",bold=True,fg=XW,bg=XB,align="center")
    xc(ws,2,8,datetime.now().strftime("%d/%m/%Y"),fg=XW,bg=XB,align="center")
    ws.row_dimensions[2].height=18
    for ci,h in enumerate(["No.","Location","QA Markup Comment","Affected Tag",
                            "Category","Sev.","Reviewer","Sheet"],1):
        xc(ws,3,ci,h,bold=True,fg="3D4BA0",bg=XL,align="center")
    ws.row_dimensions[3].height=16; ws.freeze_panes="A4"
    for ri,m in enumerate(markups,4):
        alt=XA if ri%2==0 else XW
        xc(ws,ri,1,m["id"],align="center",bg=alt)
        xc(ws,ri,2,m["location"],bg=alt,size=9)
        xc(ws,ri,3,m["comment"],bg=alt,wrap=True,size=9)
        xc(ws,ri,4,m["affected_tag"],align="center",bg=alt,size=9)
        xc(ws,ri,5,m["category"],bg=alt,size=9)
        xs(ws,ri,6,m["severity"])
        xc(ws,ri,7,m["reviewer"],bg=alt,size=9)
        xc(ws,ri,8,m["page"],align="center",bg=alt,size=9)
        ws.row_dimensions[ri].height=40 if len(m["comment"])>80 else 20

def build_excel(all_results,filename):
    wb=openpyxl.Workbook()
    build_summary(wb,all_results,filename)
    build_master_tab(wb,all_results)
    for pg in sorted(k for k,v in all_results.items() if v["total"]>0):
        r=all_results[pg]
        build_page_tab(wb,f"Sheet {pg}",r["markups"],r["drawing_info"])
    return wb



# ════════════════════════════════════════════════════════════
# UI
# ════════════════════════════════════════════════════════════
if not st.session_state.get("pdf_loaded"):
    st.markdown('<div class="sh">Upload QA-Marked Drawing PDF</div>',unsafe_allow_html=True)

pdf_file=st.file_uploader("Upload PDF",type=["pdf"],label_visibility="collapsed")


if not pdf_file:
    st.markdown("""<div style="border:2px dashed #5C6BC0;border-radius:16px;
    background:#E8EAF6;padding:40px;text-align:center;margin:12px 0">
    <div style="font-size:48px">📄</div>
    <div style="font-family:Poppins,sans-serif;font-size:18px;font-weight:600;
    color:#3D4BA0;margin-bottom:6px">Drop your QA-marked drawing PDF here</div>
    <div style="font-size:13px;color:#4A4A6A">Supports Bluebeam · Adobe Acrobat · Any annotated PDF</div>
    </div>""",unsafe_allow_html=True)
    st.markdown("<br>",unsafe_allow_html=True)
    ##c1,c2,c3=st.columns(3)
    ##c1.markdown('<div class="fcard"><div class="ficon">🔍</div><div class="ftitle">Extracts All QA Annotations</div><div class="fdesc">FreeText comments — every typed QA reviewer comment captured</div></div>',unsafe_allow_html=True)
    ##c2.markdown('<div class="fcard"><div class="ficon">🏷️</div><div class="ftitle">Auto HVAC Tag Detection</div><div class="fdesc">Automatically finds the nearest FCU, FD, OAF tag for each markup</div></div>',unsafe_allow_html=True)
    ##c3.markdown('<div class="fcard"><div class="ficon">📊</div><div class="ftitle">Professional Excel Report</div><div class="fdesc">Summary + Master sheet + Per-page tabs with severity color coding</div></div>',unsafe_allow_html=True)
    st.stop() 

if st.session_state.get("pdf_name")!=pdf_file.name:
    with st.spinner("Loading pages..."):
        pdf_bytes=pdf_file.read()
        doc=fitz.open(stream=pdf_bytes,filetype="pdf")
        pages=[]
        for page in doc:
            mat=fitz.Matrix(100/72,100/72)
            pix=page.get_pixmap(matrix=mat,alpha=False)
            img=Image.frombytes("RGB",[pix.width,pix.height],pix.samples)
            pages.append(img)
        doc.close()
        st.session_state.update({"pdf_name":pdf_file.name,"pages":pages,
                                  "pdf_bytes":pdf_bytes,"current":1,"pdf_loaded":True})


pages=st.session_state["pages"]; total_pages=len(pages)
col_img,col_side=st.columns([3,1])



with col_side:
    st.markdown(f'<div class="spanel"><h4>📋 {total_pages} Sheet(s) Loaded</h4><div style="font-size:12px;color:#4A4A6A">{pdf_file.name}</div></div>',unsafe_allow_html=True)
    st.markdown("<br>",unsafe_allow_html=True)
    if total_pages>1:
        current=st.slider("Navigate Sheets",1,total_pages,st.session_state["current"],format="Sheet %d")
        st.session_state["current"]=current
        b1,b2=st.columns(2)
        if b1.button("◀ Prev",use_container_width=True):
            st.session_state["current"]=max(1,current-1); st.rerun()
        if b2.button("Next ▶",use_container_width=True):
            st.session_state["current"]=min(total_pages,current+1); st.rerun()
    else:
        current=1
    st.markdown("<br>",unsafe_allow_html=True)
    st.info("Extracts QA annotations from **all pages** at once.")
    run_btn=st.button("🔍  Extract All QA Annotations",use_container_width=True,type="primary")



with col_img:
    st.markdown(f'<div class="sh">Drawing Preview — Sheet {current} of {total_pages}</div>',unsafe_allow_html=True)
    st.image(pages[current-1],use_container_width=True)

st.markdown("<hr>",unsafe_allow_html=True)

if run_btn:
    prog=st.progress(0); status=st.empty()
    status.info("🔍 Scanning all pages for QA annotations...")
    all_results=extract_all_markups(st.session_state["pdf_bytes"])
    prog.progress(60)
    all_m=[m for r in all_results.values() for m in r["markups"]]
    total=len(all_m)
    pages_hit=[p for p,r in all_results.items() if r["total"]>0]
    high=sum(1 for m in all_m if m["severity"]=="High")
    med=sum(1 for m in all_m if m["severity"]=="Medium")
    lo=sum(1 for m in all_m if m["severity"]=="Low")
    status.info("📊 Building Excel report...")
    wb=build_excel(all_results,pdf_file.name)
    buf=BytesIO(); wb.save(buf); buf.seek(0)
    prog.progress(100)
    status.success(f"✅ {total} QA markups extracted from {len(pages_hit)} sheet(s)!")



    st.markdown('<div class="sh">Extraction Summary</div>',unsafe_allow_html=True)
    st.markdown(f"""<div class="mrow">
      <div class="mcard blue"><div class="mval">{total}</div><div class="mlbl">Total Markups</div></div>
      <div class="mcard red"><div class="mval">{high}</div><div class="mlbl">🔴 High</div></div>
      <div class="mcard amber"><div class="mval">{med}</div><div class="mlbl">🟡 Medium</div></div>
      <div class="mcard green"><div class="mval">{lo}</div><div class="mlbl">🟢 Low</div></div>
      <div class="mcard gray"><div class="mval">{len(pages_hit)}</div><div class="mlbl">Sheets</div></div>
    </div>""",unsafe_allow_html=True)

    page_counts={pg:all_results[pg]["total"] for pg in pages_hit}
    df_pg=pd.DataFrame({"Sheet":[f"S{p}" for p in page_counts.keys()],"Markups":list(page_counts.values())})
    st.markdown('<div class="sh">Markups per Sheet</div>',unsafe_allow_html=True)
    st.bar_chart(df_pg.set_index("Sheet"))
    st.markdown("<hr>",unsafe_allow_html=True)

    st.markdown('<div class="sh">All QA Annotations</div>',unsafe_allow_html=True)
    df=pd.DataFrame(all_m).rename(columns={"id":"No.","page":"Sheet",
        "drawing_number":"Drawing No","location":"Location","comment":"QA Comment",
        "affected_tag":"Tag","category":"Category","severity":"Severity","reviewer":"Reviewer"})
    fc1,fc2,fc3=st.columns(3)
    sev_f=fc1.multiselect("Severity",["High","Medium","Low"],default=["High","Medium","Low"])
    pg_f=fc2.multiselect("Sheet",sorted(df["Sheet"].unique()),default=sorted(df["Sheet"].unique()))
    cat_f=fc3.multiselect("Category",sorted(df["Category"].unique()),default=sorted(df["Category"].unique()))
    view=df[df["Severity"].isin(sev_f)&df["Sheet"].isin(pg_f)&df["Category"].isin(cat_f)]
    st.caption(f"Showing **{len(view)}** of **{total}** markups")
    st.dataframe(view[["No.","Sheet","Location","QA Comment","Tag","Category","Severity","Reviewer"]],
        use_container_width=True,hide_index=True,
        column_config={"No.":st.column_config.NumberColumn(width="small"),
                       "Sheet":st.column_config.NumberColumn(width="small"),
                       "Severity":st.column_config.TextColumn(width="small"),
                       "QA Comment":st.column_config.TextColumn(width="large")})
    st.markdown("<hr>",unsafe_allow_html=True)
    fn=f"DrawCheck_QA_{os.path.splitext(pdf_file.name)[0]}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    st.download_button("📥  Download Complete QA Report (.xlsx)",
        data=buf.getvalue(),file_name=fn,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,type="primary")



st.markdown(f"""<hr><div style="text-align:center;color:#4A4A6A;font-size:12px;padding:8px 0">
DrawCheck AI v1.0 &nbsp;|&nbsp; SG Design Nepal &nbsp;|&nbsp;
Extracts QA FreeText Annotations &nbsp;|&nbsp; No API required</div>""",unsafe_allow_html=True)
